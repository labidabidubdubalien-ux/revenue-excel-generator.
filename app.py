import io
import re
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
import streamlit as st

# Config Halaman Website
st.set_page_config(
    page_title="Revenue Converter & Summarizer", page_icon="📊", layout="wide"
)


# ==========================================
# 1. LOGIKA KATEGORI & PARSER
# ==========================================
def get_category(description):
    desc_lower = description.lower()
    if "sugar & spice" in desc_lower:
        return "Sugar & Spice"
    elif "lobby lounge" in desc_lower:
        return "Lobby Lounge"
    elif "room service" in desc_lower:
        return "Room Service"
    elif "minibar" in desc_lower:
        return "Minibar"
    elif "banquet" in desc_lower:
        return "Banquet"
    elif "spa" in desc_lower:
        return "Spa"
    elif "laundry" in desc_lower:
        return "Guest Laundry"
    elif any(
        kw in desc_lower
        for kw in ["accomodation ++", "Upsell Late CheckOut ++", "Upselling", "Accomodation Nett", 
                   "Upsell Early Check In", "Upsell Early Check In - Rebate"]
    ):
        return "Room Revenue"
    else:
        return "Other Revenue"


def parse_and_categorize(text):
    lines = [
        line.strip()
        for line in text.strip().split("\n")
        if line.strip() and line.strip().lower() != "revenue"
    ]
    descriptions, amounts = [], []
    amount_pattern = re.compile(r"^-?\s*[\d,]+$")

    for line in lines:
        if amount_pattern.match(line):
            clean_num = line.replace(" ", "").replace(",", "")
            amounts.append(int(clean_num))
        else:
            descriptions.append(line)

    categorized_data = []
    category_totals = {}

    for desc, amount in zip(descriptions, amounts):
        cat = get_category(desc)
        categorized_data.append((cat, desc, amount))
        category_totals[cat] = category_totals.get(cat, 0) + amount

    return categorized_data, category_totals


# ==========================================
# 2. GENERATE EXCEL (SUMMARY & DETAIL SHEET)
# ==========================================
def generate_excel(revenue_data, category_totals):
    wb = openpyxl.Workbook()

    NAVY, ICE_BLUE, BORDER_GRAY, ZEBRA_FILL = (
        "1B365D",
        "F2F5F9",
        "D9D9D9",
        "F9FAFC",
    )
    font_title = Font(name="Calibri", size=16, bold=True, color=NAVY)
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="555555")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_regular = Font(name="Calibri", size=11)
    font_total = Font(name="Calibri", size=12, bold=True, color=NAVY)

    fill_header = PatternFill(
        start_color=NAVY, end_color=NAVY, fill_type="solid"
    )
    fill_zebra = PatternFill(
        start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid"
    )
    fill_total = PatternFill(
        start_color=ICE_BLUE, end_color=ICE_BLUE, fill_type="solid"
    )

    thin_border = Border(
        left=Side(style="thin", color=BORDER_GRAY),
        right=Side(style="thin", color=BORDER_GRAY),
        top=Side(style="thin", color=BORDER_GRAY),
        bottom=Side(style="thin", color=BORDER_GRAY),
    )
    double_bottom_border = Border(
        top=Side(style="thin", color=NAVY),
        bottom=Side(style="double", color=NAVY),
        left=Side(style="thin", color=BORDER_GRAY),
        right=Side(style="thin", color=BORDER_GRAY),
    )

    num_format = '#,##0;(#,##0);"-";@'

    # SHEET 1: SUMMARY
    ws_summary = wb.active
    ws_summary.title = "Category Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary.merge_cells("A1:B1")
    ws_summary["A1"] = "InterContinental Jakarta Pondok Indah"
    ws_summary["A1"].font = font_title

    ws_summary.merge_cells("A2:B2")
    ws_summary["A2"] = "Revenue Summary - Per Category Subtotal"
    ws_summary["A2"].font = font_subtitle

    summary_headers = ["Category", "Total Amount (IDR)"]
    for col_idx, header in enumerate(summary_headers, 1):
        cell = ws_summary.cell(row=4, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(
            horizontal="left" if col_idx == 1 else "right", vertical="center"
        )
        cell.border = thin_border

    ws_summary.row_dimensions[4].height = 25

    start_row_sum = 5
    for i, (cat_name, total_val) in enumerate(category_totals.items()):
        row_num = start_row_sum + i
        c_cat = ws_summary.cell(row=row_num, column=1, value=cat_name)
        c_val = ws_summary.cell(row=row_num, column=2, value=total_val)

        c_cat.font, c_val.font = font_regular, font_regular
        c_val.number_format = num_format
        c_cat.alignment = Alignment(horizontal="left")
        c_val.alignment = Alignment(horizontal="right")

        if i % 2 == 1:
            c_cat.fill, c_val.fill = fill_zebra, fill_zebra

        c_cat.border, c_val.border = thin_border, thin_border

    total_row_sum = start_row_sum + len(category_totals)
    ws_summary.cell(
        row=total_row_sum, column=1, value="Grand Total"
    ).font = font_total
    ws_summary.cell(row=total_row_sum, column=1).fill = fill_total
    ws_summary.cell(row=total_row_sum, column=1).alignment = Alignment(
        horizontal="right"
    )
    ws_summary.cell(row=total_row_sum, column=1).border = double_bottom_border

    c_tot_val = ws_summary.cell(
        row=total_row_sum,
        column=2,
        value=f"=SUM(B{start_row_sum}:B{total_row_sum-1})",
    )
    c_tot_val.font = font_total
    c_tot_val.fill = fill_total
    c_tot_val.alignment = Alignment(horizontal="right")
    c_tot_val.number_format = num_format
    c_tot_val.border = double_bottom_border

    ws_summary.column_dimensions["A"].width = 30
    ws_summary.column_dimensions["B"].width = 25

    # SHEET 2: DETAIL
    ws_detail = wb.create_sheet(title="Transaction Detail")
    ws_detail.views.sheetView[0].showGridLines = True
    ws_detail.freeze_panes = "A5"

    ws_detail.merge_cells("A1:C1")
    ws_detail["A1"] = "InterContinental Jakarta Pondok Indah"
    ws_detail["A1"].font = font_title

    ws_detail.merge_cells("A2:C2")
    ws_detail["A2"] = "Trial Balance - Categorized Detail Section"
    ws_detail["A2"].font = font_subtitle

    headers = ["Category", "Account Description", "Amount (IDR)"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws_detail.cell(row=4, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(
            horizontal=(
                "center"
                if col_idx == 1
                else ("left" if col_idx == 2 else "right")
            ),
            vertical="center",
        )
        cell.border = thin_border

    ws_detail.row_dimensions[4].height = 25

    start_row_det = 5
    for i, (category, desc, amount) in enumerate(revenue_data):
        row_num = start_row_det + i
        cell_cat = ws_detail.cell(row=row_num, column=1, value=category)
        cell_desc = ws_detail.cell(row=row_num, column=2, value=desc)
        cell_amount = ws_detail.cell(row=row_num, column=3, value=amount)

        cell_cat.alignment = Alignment(horizontal="center")
        cell_desc.alignment = Alignment(horizontal="left")
        cell_amount.alignment = Alignment(horizontal="right")

        cell_cat.font, cell_desc.font, cell_amount.font = (
            font_regular,
            font_regular,
            font_regular,
        )
        cell_amount.number_format = num_format

        if i % 2 == 1:
            cell_cat.fill, cell_desc.fill, cell_amount.fill = (
                fill_zebra,
                fill_zebra,
                fill_zebra,
            )

        cell_cat.border, cell_desc.border, cell_amount.border = (
            thin_border,
            thin_border,
            thin_border,
        )

    total_row_det = start_row_det + len(revenue_data)
    ws_detail.cell(row=total_row_det, column=1, value="").border = (
        double_bottom_border
    )

    cell_label = ws_detail.cell(
        row=total_row_det, column=2, value="Revenue Total"
    )
    cell_label.font = font_total
    cell_label.fill = fill_total
    cell_label.alignment = Alignment(horizontal="right")
    cell_label.border = double_bottom_border

    cell_value = ws_detail.cell(
        row=total_row_det,
        column=3,
        value=f"=SUM(C{start_row_det}:C{total_row_det-1})",
    )
    cell_value.font = font_total
    cell_value.fill = fill_total
    cell_value.alignment = Alignment(horizontal="right")
    cell_value.number_format = num_format
    cell_value.border = double_bottom_border

    ws_detail.column_dimensions["A"].width = 20
    ws_detail.column_dimensions["B"].width = 48
    ws_detail.column_dimensions["C"].width = 22

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ==========================================
# 3. ANTARMUKA WEBSITE (STREAMLIT UI)
# ==========================================
st.title("📊 Trial Balance Revenue Converter & Summarizer")
st.write(
    "Paste data mentah di bawah untuk memproses kategori, melihat ringkasan, dan mengunduh laporan Excel."
)

raw_input = st.text_area("Paste Raw Data di Sini:", height=200)

if st.button("🚀 Proses & Hitung Ringkasan"):
    if raw_input.strip():
        data_detail, data_summary = parse_and_categorize(raw_input)
        excel_file = generate_excel(data_detail, data_summary)

        st.success(
            f"Berhasil memproses {len(data_detail)} item ke dalam {len(data_summary)} kategori!"
        )

        # Tombol Download Excel
        st.download_button(
            label="📥 Download File Excel (Ringkasan + Detail)",
            data=excel_file,
            file_name="Trial_Balance_Revenue_Report.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        st.markdown("---")

        # BARIS 1: RINGKASAN & GRAFIK
        col1, col2 = st.columns([1, 1])

        with col1:
            st.subheader("📋 Ringkasan Total per Kategori")
            df_summary = pd.DataFrame(
                list(data_summary.items()), columns=["Kategori", "Total (IDR)"]
            )
            df_summary_formatted = df_summary.copy()
            df_summary_formatted["Total (IDR)"] = df_summary_formatted[
                "Total (IDR)"
            ].apply(lambda x: f"{x:,.0f}")
            st.dataframe(
                df_summary_formatted, use_container_width=True, hide_index=True
            )

        with col2:
            st.subheader("📈 Grafik Kontribusi Kategori")
            st.bar_chart(df_summary.set_index("Kategori"))

        st.markdown("---")

        # BARIS 2: TABEL DETAIL TRANSAKSI
        st.subheader("📑 Detail Transaksi Lengkap")

        df_detail = pd.DataFrame(
            data_detail,
            columns=["Category", "Account Description", "Amount (IDR)"],
        )
        df_detail_formatted = df_detail.copy()
        df_detail_formatted["Amount (IDR)"] = df_detail_formatted[
            "Amount (IDR)"
        ].apply(lambda x: f"{x:,.0f}")

        # Tampilan Tabel Detail
        st.dataframe(
            df_detail_formatted, use_container_width=True, hide_index=True
        )

    else:
        st.warning("Silakan paste data mentah terlebih dahulu.")
